"""comparison: 対比。mode = two-column | pros-cons | table。

data 契約: slide-expression/references/comparison.md
"""

from __future__ import annotations

from pptx.enum.text import PP_ALIGN

from deckgen import layout


def render(pslide, theme, slide, region):
    data = slide.get("data") or {}
    mode = data.get("mode", "two-column")
    if mode == "table":
        _table(pslide, theme, data, region)
    else:
        _two_column(pslide, theme, data, region, mode)


def _two_column(pslide, theme, data, region, mode):
    left, top, width, height = region
    pros_cons = mode == "pros-cons"
    gap = layout.SPACE_4
    col_w = (width - gap) // 2
    note = data.get("note")
    # note は region 内に収める（本文 = height − 注記行 − 間隔）
    note_h = layout.NOTE_H if note else 0
    body_h = height - note_h - (layout.NOTE_GAP if note else 0)

    # 見出し・アクセントバーの色は意味色で統一する:
    # pros-cons は good/bad、通常比較は accent2/accent（muted は「対比」を表さない）。
    cols = [
        ("left", data.get("left") or {}, theme["good"] if pros_cons else theme["accent2"]),
        ("right", data.get("right") or {}, theme["bad"] if pros_cons else theme["accent"]),
    ]
    label_top = top + layout.CARD_PAD_Y
    items_top = label_top + layout.CARD_LABEL_H + layout.SPACE_1
    for i, (_key, col, head_color) in enumerate(cols):
        x = left + i * (col_w + gap)
        layout.add_card(pslide, x, top, col_w, body_h, theme,
                        accent_bar=head_color)
        # ラベル見出し（バー分だけ左余白を追加）
        layout.add_textbox(
            pslide, x + layout.CARD_TEXT_LEFT, label_top,
            col_w - layout.CARD_TEXT_LEFT - layout.CARD_PAD_X, layout.CARD_LABEL_H,
            col.get("label", ""), size=layout.FONT_LEAD, color=head_color, bold=True,
        )
        # 項目
        items = [str(v) for v in (col.get("items") or [])]
        if items:
            layout.add_bullets(
                pslide, x + layout.CARD_TEXT_LEFT, items_top,
                col_w - layout.CARD_TEXT_LEFT - layout.CARD_PAD_X,
                top + body_h - items_top - layout.CARD_PAD_Y, items,
                size=layout.FONT_BODY, color=theme["fg"],
                line_spacing=1.25, space_after=6, autofit=True,
            )
    if note:
        layout.add_textbox(
            pslide, left, top + body_h + layout.NOTE_GAP, width, note_h,
            str(note), size=layout.FONT_CAPTION, color=theme["muted"],
        )


def _table(pslide, theme, data, region):
    left, top, width, height = region
    axes = [str(a) for a in (data.get("axes") or [])]
    columns = data.get("columns") or []
    if not columns:
        return
    n_rows = 1 + len(axes)
    n_cols = 1 + len(columns)
    # 高さは行数に合わせて上詰め（最大 region 内）
    row_h = min(layout.TABLE_ROW_MAX_H, height // max(n_rows, 1))
    table_h = row_h * n_rows
    table = layout.add_table(pslide, left, top, width, table_h, n_rows, n_cols)

    # 1列目を細め、データ列を等幅に
    first_w = layout.TABLE_FIRST_COL_W
    rest = (width - first_w) // len(columns)
    table.columns[0].width = first_w
    for c in range(1, n_cols):
        table.columns[c].width = rest

    # ヘッダ行: [空, 列名...]
    layout.style_cell(table.cell(0, 0), "", color=theme["on_accent"],
                      fill=theme["accent"])
    for c, col in enumerate(columns, start=1):
        layout.style_cell(
            table.cell(0, c), str(col.get("name", "")),
            size=layout.FONT_BODY, color=theme["on_accent"], bold=True,
            fill=theme["accent"], align=PP_ALIGN.CENTER,
        )
    # 各評価軸の行
    for r, axis in enumerate(axes, start=1):
        zebra = theme["card"] if r % 2 == 1 else theme["bg"]
        layout.style_cell(
            table.cell(r, 0), axis, size=layout.FONT_SMALL, color=theme["fg"],
            bold=True, fill=theme["card"],
        )
        for c, col in enumerate(columns, start=1):
            vals = col.get("values") or []
            text = str(vals[r - 1]) if r - 1 < len(vals) else ""
            layout.style_cell(
                table.cell(r, c), text, size=layout.FONT_SMALL, color=theme["fg"],
                fill=zebra, align=PP_ALIGN.CENTER,
            )
